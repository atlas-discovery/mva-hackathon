"""Extract HPO terms from a phenotype document (.docx, .txt, .md)."""

from __future__ import annotations

import re
from pathlib import Path

HPO_ID = re.compile(r"HP:\d{7}")


def read_text(path: str) -> str:
    p = Path(path)
    if p.suffix.lower() == ".docx":
        try:
            import docx
        except ImportError as e:
            raise RuntimeError("reading .docx needs: pip install 'mva-hackathon[docx]'") from e
        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return p.read_text()


def hpo_terms(path: str) -> list[dict]:
    """Return unique HPO IDs with the text of the line they appeared on."""
    seen: dict[str, str] = {}
    for line in read_text(path).splitlines():
        for hid in HPO_ID.findall(line):
            seen.setdefault(hid, line.strip()[:160])
    return [{"id": k, "context": v} for k, v in seen.items()]
